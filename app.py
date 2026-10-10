# importar as bibliotecas necessárias
import streamlit as st
from PyPDF2 import PdfReader
import docx
from docx.shared import Pt
import os
from dotenv import load_dotenv
import google.generativeai as genai
import io
import re

# --- CONFIGURAÇÃO DE SEGURANÇA E IA ---
load_dotenv()
chave_api = os.getenv("GEMINI_API_KEY") 
genai.configure(api_key=chave_api)
modelo_ia = genai.GenerativeModel("gemini-flash-latest")

st.set_page_config(page_title="Gerador de Modelos Inteligente", page_icon="📄", layout="centered")

# --- GERENCIAMENTO DE MEMÓRIA (SESSION STATE) ---
if "texto_modelo" not in st.session_state:
    st.session_state.texto_modelo = None
if "variaveis_encontradas" not in st.session_state:
    st.session_state.variaveis_encontradas = []

# --- FUNÇÕES DE EXTRAÇÃO DE TEXTO ---
def extrair_texto_pdf(arquivo_pdf):
    reader = PdfReader(arquivo_pdf)
    texto = ""
    for page in reader.pages:
        texto += page.extract_text() + "\n"
    return texto

def extrair_texto_docx(arquivo_docx):
    doc = docx.Document(arquivo_docx)
    texto = ""
    for paragraph in doc.paragraphs:
        texto += paragraph.text + "\n"
    return texto

def ler_arquivo(arquivo_upload):
    if arquivo_upload.type == "application/pdf":
        return extrair_texto_pdf(arquivo_upload)
    elif arquivo_upload.type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        return extrair_texto_docx(arquivo_upload)
    else:
        return ""

# --- MOTOR DE COMPARAÇÃO DO GEMINI ---
def comparar_documentos_com_gemini(lista_de_textos):
    textos_unidos = ""
    for i, texto in enumerate(lista_de_textos):
        textos_unidos += f"\n\n--- DOCUMENTO {i+1} ---\n{texto}"

    prompt = f"""
    Você é um especialista jurídico e analista de contratos.
    Abaixo, forneço o texto de 2 ou mais documentos semelhantes.
    
    Sua missão é comparar estes documentos e gerar um ÚNICO "Documento Modelo" (Template).
    
    REGRA 1: Tudo o que for texto idêntico ou comum entre os documentos fornecidos deve ser MANTIDO INTACTO.
    REGRA 2: Tudo o que variar entre os documentos (como Nomes, CNPJ, CPF, RG, Endereços, Valores Numéricos, Taxas, Prazos, Garantias, Datas) deve ser substituído por campos indicativos no formato: [NOME_DA_VARIAVEL].
    REGRA 3: Retorne EXCLUSIVAMENTE o texto final do Documento Modelo. Não inclua saudações, não explique o que fez, apenas entregue o texto pronto para uso.
    
    Textos dos Documentos:
    {textos_unidos}
    """
    
    resposta = modelo_ia.generate_content(prompt)
    return resposta.text

# --- GERAÇÃO DO WORD COM FONTE PERSONALIZADA ---
def gerar_documento_word(texto_para_word, eh_template=True):
    doc = docx.Document()
    
    estilo_normal = doc.styles['Normal']
    fonte = estilo_normal.font
    fonte.name = 'Arial'
    fonte.size = Pt(11)
    
    if eh_template:
        doc.add_heading('Modelo de Documento Gerado por IA', level=1)
        doc.add_paragraph('Atenção: Os campos que variam foram marcados com colchetes.\n')
    else:
        doc.add_heading('Documento Final Preenchido', level=1)

    doc.add_paragraph(texto_para_word)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

# --- INTERFACE DO USUÁRIO ---
st.title("Gerador de Modelos Inteligente 🤖")
st.write("Faça o upload de documentos semelhantes. A IA criará um Template inteligente, que você poderá baixar ou preencher diretamente aqui no aplicativo!")

arquivos_upload = st.file_uploader(
    label="Selecione os documentos base:", 
    type=["pdf", "docx"],
    accept_multiple_files=True 
)

st.divider()

# ETAPA 1: GERAÇÃO DO TEMPLATE PELA IA
if st.button("Analisar e Gerar Documento Modelo"):
    if arquivos_upload and len(arquivos_upload) >= 2:
        with st.spinner("Lendo e comparando os documentos..."):
            textos_extraidos = [ler_arquivo(arq) for arq in arquivos_upload if ler_arquivo(arq)]
            
            if textos_extraidos:
                texto_gerado = comparar_documentos_com_gemini(textos_extraidos)
                
                # --- NOVIDADE: NORMALIZAÇÃO DE COLCHETES ---
                # Substituímos os duplos por simples para criar um padrão único
                texto_gerado = texto_gerado.replace("[[", "[").replace("]]", "]")
                
                st.session_state.texto_modelo = texto_gerado
                
                # --- NOVIDADE: ATUALIZAÇÃO DO DETETIVE (REGEX) ---
                # Agora buscamos apenas por um colchete simples de cada lado: \[ e \]
                variaveis = list(set(re.findall(r'\[(.*?)\]', texto_gerado)))
                st.session_state.variaveis_encontradas = variaveis
                
                st.success("Template gerado com sucesso! Veja as opções abaixo.")
    else:
        st.warning("⚠️ Envie pelo menos dois documentos para comparação.")

# ETAPA 2: EXIBIÇÃO DAS OPÇÕES
if st.session_state.texto_modelo:
    st.divider()
    
    st.subheader("Opção A: Baixar Template em Branco")
    arquivo_template = gerar_documento_word(st.session_state.texto_modelo, eh_template=True)
    st.download_button(
        label="⬇️ Baixar Template (.docx)",
        data=arquivo_template,
        file_name="Template_Em_Branco.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    
    st.divider()
    
    st.subheader("Opção B: Preencher Dados para o Documento Final")
    st.write("Preencha os campos abaixo encontrados no seu modelo:")
    
    with st.form("form_preenchimento"):
        respostas_usuario = {}
        
        for var in st.session_state.variaveis_encontradas:
            respostas_usuario[var] = st.text_input(label=f"Digite o valor para: {var}")
            
        botao_gerar_final = st.form_submit_button("Gerar Documento Preenchido")
        
    if botao_gerar_final:
        texto_final = st.session_state.texto_modelo
        
        # --- NOVIDADE: SUBSTITUIÇÃO COM COLCHETE SIMPLES ---
        for var, resposta in respostas_usuario.items():
            if resposta: 
                # Agora substituímos usando [var] em vez de [[var]]
                texto_final = texto_final.replace(f"[{var}]", resposta)
        
        with st.spinner("Gerando arquivo final..."):
            arquivo_preenchido = gerar_documento_word(texto_final, eh_template=False)
            
        st.success("Documento final gerado com sucesso!")
        
        st.download_button(
            label="⬇️ Baixar Documento PREENCHIDO (.docx)",
            data=arquivo_preenchido,
            file_name="Documento_Final_Preenchido.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ) 